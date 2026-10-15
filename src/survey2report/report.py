"""Build a professional .pptx report from the analysis and its charts."""

from __future__ import annotations

from collections.abc import Sequence
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from survey2report.charts import Chart
from survey2report.cleaninglog import CleaningLog
from survey2report.slides import (
    ACCENT,
    BODY,
    CONTENT_WIDTH,
    INK,
    MARGIN,
    MUTED,
    SLIDE_HEIGHT,
    SLIDE_WIDTH,
    add_bullets,
    add_footer,
    add_table,
    add_text,
    add_title,
    new_slide,
)
from survey2report.stats import AnalysisResult

MAX_TABLE_ROWS = 10
CHART_MAX_HEIGHT = Inches(4.3)


def humanise(stem: str) -> str:
    """Turn a file stem into a readable slide title.

    Args:
        stem: File name without its extension.

    Returns:
        A title-cased name that keeps short acronyms upper case.
    """
    words = stem.replace("_", " ").replace("-", " ").split()
    return " ".join(
        word.upper() if word.isupper() and len(word) > 1 else word.capitalize() for word in words
    )


def _title_slide(prs: Presentation, analysis: AnalysisResult, generated_on: str) -> None:
    """Add the cover slide."""
    slide = new_slide(prs)
    add_text(slide, "SURVEY2REPORT", MARGIN, Inches(2.1), CONTENT_WIDTH, Inches(0.4),
             size=13, bold=True, color=ACCENT)
    add_text(slide, humanise(Path(analysis.filename).stem), MARGIN, Inches(2.6), CONTENT_WIDTH,
             Inches(1.1), size=40, bold=True, color=INK)
    add_text(slide, f"Automated findings from {analysis.rows} cleaned responses", MARGIN,
             Inches(3.8), CONTENT_WIDTH, Inches(0.5), size=17, color=BODY)
    add_text(slide, f"{analysis.column_count} analysed columns  ·  generated {generated_on}",
             MARGIN, Inches(4.4), CONTENT_WIDTH, Inches(0.4), size=12, color=MUTED)


def _findings_slide(prs: Presentation, analysis: AnalysisResult) -> None:
    """Add the key findings slide."""
    slide = new_slide(prs)
    add_title(slide, "Key findings", "What the data says")
    add_bullets(slide, analysis.findings, size=16)


def _chart_slide(prs: Presentation, chart: Chart) -> None:
    """Add one chart slide with its caption."""
    slide = new_slide(prs)
    add_title(slide, chart.title, "Evidence")
    picture = slide.shapes.add_picture(str(chart.path), MARGIN, Inches(1.7), width=CONTENT_WIDTH)
    if picture.height > CHART_MAX_HEIGHT:
        picture.width = int(CHART_MAX_HEIGHT * picture.width / picture.height)
        picture.height = CHART_MAX_HEIGHT
        picture.left = int((SLIDE_WIDTH - picture.width) / 2)
    add_text(slide, chart.caption, MARGIN, SLIDE_HEIGHT - Inches(1.45), CONTENT_WIDTH,
             Inches(0.7), size=12, color=MUTED)


def _methodology_slide(prs: Presentation, analysis: AnalysisResult, log: CleaningLog) -> None:
    """Add the methodology and data quality slide."""
    slide = new_slide(prs)
    add_title(slide, "Methodology & data quality", "How this was produced")
    bullets = [
        "Each file is read as raw text; every column is typed as numeric, date, "
        "categorical or free text.",
        "Answers are repaired in place; each change is recorded with a count in "
        "cleaning_log.md.",
        *log.quality_summary()[:3],
        f"Charts and figures come from the cleaned table: {len(analysis.numeric)} numeric, "
        f"{len(analysis.categorical)} categorical and {len(analysis.dates)} date columns.",
    ]
    add_bullets(slide, bullets, size=14)


def _metrics_appendix(prs: Presentation, analysis: AnalysisResult) -> None:
    """Add the numeric metrics appendix table."""
    slide = new_slide(prs)
    add_title(slide, "Appendix: metric detail", "Summary statistics")
    rows = [
        [
            summary.column,
            f"{summary.count}",
            f"{summary.missing}",
            f"{summary.mean:,.2f}",
            f"{summary.median:,.2f}",
            f"{summary.minimum:,.2f}",
            f"{summary.maximum:,.2f}",
        ]
        for summary in analysis.numeric[:MAX_TABLE_ROWS]
    ]
    headers = ["Column", "N", "Missing", "Mean", "Median", "Min", "Max"]
    if rows:
        add_table(slide, headers, rows)
    else:
        add_bullets(slide, ["No numeric columns were found in this dataset."])


def _categories_appendix(prs: Presentation, analysis: AnalysisResult) -> None:
    """Add the categorical detail appendix table."""
    slide = new_slide(prs)
    add_title(slide, "Appendix: category detail", "Most common answers")
    rows = [
        [
            summary.column,
            f"{summary.unique}",
            summary.top,
            f"{summary.top_share:.0%}",
            ", ".join(f"{level} ({count})" for level, count in summary.distribution[1:4]) or "—",
        ]
        for summary in analysis.categorical[:MAX_TABLE_ROWS]
    ]
    headers = ["Column", "Levels", "Most common", "Share", "Also common"]
    if rows:
        add_table(slide, headers, rows)
    else:
        add_bullets(slide, ["No categorical columns were found in this dataset."])


def _cleaning_appendix(prs: Presentation, log: CleaningLog) -> None:
    """Add the cleaning actions appendix table."""
    slide = new_slide(prs)
    shown = log.entries[:MAX_TABLE_ROWS]
    kicker = f"{len(log.entries)} recorded actions"
    if len(shown) < len(log.entries):
        kicker = f"showing {len(shown)} of {len(log.entries)} recorded actions"
    add_title(slide, "Appendix: cleaning log", kicker)
    rows = [
        [entry.column or "whole table", entry.issue, entry.action, str(entry.count)]
        for entry in shown
    ]
    headers = ["Column", "Issue found", "Action taken", "Cells"]
    if rows:
        add_table(slide, headers, rows)
    else:
        add_bullets(slide, ["Nothing needed cleaning in this dataset."])


def build_report(
    analysis: AnalysisResult,
    charts: Sequence[Chart],
    log: CleaningLog,
    path: str | Path,
    generated_on: str | None = None,
) -> Path:
    """Assemble the full PowerPoint report and write it to disk.

    Args:
        analysis: Analysis result for one dataset.
        charts: Rendered charts, in the order they should appear.
        log: Cleaning log for the same dataset.
        path: Destination .pptx file path.
        generated_on: Cover date override, mainly for reproducible tests.

    Returns:
        The path that was written.
    """
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    _title_slide(prs, analysis, generated_on or date.today().strftime("%d %B %Y"))
    _findings_slide(prs, analysis)
    for chart in charts:
        _chart_slide(prs, chart)
    _methodology_slide(prs, analysis, log)
    _metrics_appendix(prs, analysis)
    _categories_appendix(prs, analysis)
    _cleaning_appendix(prs, log)

    label = f"{analysis.filename}  ·  survey2report"
    for page, slide in enumerate(prs.slides, start=1):
        if page > 1:
            add_footer(slide, label, page)
    prs.save(target)
    return target
