"""Low-level slide building blocks: palette, text, bullets, tables and footers."""

from __future__ import annotations

from collections.abc import Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.text.text import TextFrame
from pptx.util import Inches, Pt

INK = RGBColor(0x1F, 0x2A, 0x37)
BODY = RGBColor(0x44, 0x4E, 0x5A)
ACCENT = RGBColor(0x3F, 0x6F, 0x9F)
MUTED = RGBColor(0x8B, 0x98, 0xA5)
FONT = "Calibri"

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_WIDTH = SLIDE_WIDTH - 2 * MARGIN
BODY_TOP = Inches(1.8)


def new_slide(prs: Presentation) -> Slide:
    """Add a blank slide to the deck."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_text(
    slide: Slide,
    text: str,
    left: int,
    top: int,
    width: int,
    height: int,
    *,
    size: int = 14,
    bold: bool = False,
    color: RGBColor = BODY,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> TextFrame:
    """Add a single-run text box and return its text frame."""
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    paragraph.font.name = FONT
    return frame


def add_title(slide: Slide, title: str, kicker: str | None = None) -> None:
    """Draw the standard heading with an optional kicker above it."""
    if kicker:
        add_text(slide, kicker.upper(), MARGIN, Inches(0.45), CONTENT_WIDTH, Inches(0.3),
                 size=11, bold=True, color=ACCENT)
    add_text(slide, title, MARGIN, Inches(0.8), CONTENT_WIDTH, Inches(0.7),
             size=26, bold=True, color=INK)


def add_bullets(
    slide: Slide, items: Sequence[str], top: int = BODY_TOP, size: int = 15
) -> None:
    """Add a bulleted list filling the body of the slide."""
    frame = add_text(slide, "", MARGIN, top, CONTENT_WIDTH, SLIDE_HEIGHT - top - Inches(0.8))
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = BODY
        paragraph.font.name = FONT
        paragraph.space_after = Pt(14)


def add_table(slide: Slide, headers: Sequence[str], rows: Sequence[Sequence[str]]) -> None:
    """Add a data table under the slide heading."""
    shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), MARGIN, BODY_TOP, CONTENT_WIDTH, Inches(0.4)
    )
    table = shape.table
    for column, header in enumerate(headers):
        cell = table.cell(0, column)
        cell.text = str(header)
        run = cell.text_frame.paragraphs[0]
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = INK
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            cell = table.cell(row_index, column)
            cell.text = str(value)
            cell.text_frame.paragraphs[0].font.size = Pt(10)
            cell.text_frame.paragraphs[0].font.color.rgb = BODY


def add_footer(slide: Slide, label: str, page: int) -> None:
    """Add the dataset label and page number to the bottom of the slide."""
    add_text(slide, label, MARGIN, SLIDE_HEIGHT - Inches(0.55), Inches(9), Inches(0.3),
             size=9, color=MUTED)
    add_text(slide, str(page), SLIDE_WIDTH - MARGIN - Inches(0.8), SLIDE_HEIGHT - Inches(0.55),
             Inches(0.8), Inches(0.3), size=9, color=MUTED, align=PP_ALIGN.RIGHT)
